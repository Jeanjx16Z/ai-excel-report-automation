"""Actual synthetic workbook upload-to-PDF contract, security, and cleanup tests."""

from contextlib import ExitStack
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event, Thread
import json
import unittest
from unittest.mock import patch
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader

from mining_api.app import create_app
from mining_api.config import Settings
from mining_api.errors import ApiError
from mining_api.pdf_renderer import render_pdf
from mining_api.report_service import XLSX_TYPE
from mining_api.workbook_safety import MAX_UPLOAD_BYTES

ROOT = Path(__file__).resolve().parents[3]
FAKE_SECRET = "fake-local-slice-unit-test-value-000000000"
ORIGIN = "http://localhost:3000"


def workbook_bytes(invalid=False, formula=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["date", "department", "category", "budget", "actual_cost"])
    sheet.append(["2025-01-01", "Mining Operations", "Diesel Fuel", 100, "=1+1" if formula else 120])
    sheet.append(["2025-01-02", "HSE", "Safety Equipment", 200, 180])
    sheet.append(["2025-01-03", "Logistics", "Hauling Transportation", 300, 300])
    if invalid:
        sheet["C1"] = "unknown"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def modify_archive(data, replacements=None, extra=None):
    destination = BytesIO()
    with ZipFile(BytesIO(data)) as archive, ZipFile(destination, "w", ZIP_DEFLATED) as output:
        for name in archive.namelist():
            value = archive.read(name)
            if replacements and name in replacements:
                value = replacements[name](value)
            output.writestr(name, value)
        for name, value in (extra or {}).items():
            output.writestr(name, value)
    return destination.getvalue()


class LocalSliceTests(unittest.TestCase):
    def setUp(self):
        self.now = datetime(2026, 10, 31, tzinfo=timezone.utc)
        (ROOT / ".runtime").mkdir(exist_ok=True)
        self.temp = TemporaryDirectory(dir=ROOT / ".runtime")
        self.settings = Settings(proxy_secret=FAKE_SECRET, cookie_secure=False, runtime_dir=Path(self.temp.name))
        self.app = create_app(self.settings, clock=lambda: self.now)
        self.client = self.new_client()
        self.csrf = self.bootstrap(self.client)
        self.service = self.app.state.reports
        self.guards = ExitStack()
        self.guards.enter_context(patch("dotenv.load_dotenv", side_effect=AssertionError("No dotenv in tests")))
        self.guards.enter_context(patch("httpx.HTTPTransport.handle_request", side_effect=AssertionError("No outbound HTTP")))
        self.guards.enter_context(patch("httpx.AsyncHTTPTransport.handle_async_request", side_effect=AssertionError("No outbound HTTP")))

    def tearDown(self):
        self.guards.close()
        self.client.close()
        self.temp.cleanup()

    def new_client(self):
        return TestClient(self.app, base_url="http://localhost", raise_server_exceptions=False,
                          headers={"x-api-proxy-secret": FAKE_SECRET, "origin": ORIGIN})

    def bootstrap(self, client):
        response = client.post("/api/v1/session", json={})
        self.assertEqual(201, response.status_code)
        return response.json()["csrf_token"]

    def intent(self, data, filename="synthetic.xlsx"):
        response = self.client.post("/api/v1/uploads/intents", json={"filename": filename,
            "size_bytes": len(data), "content_type": XLSX_TYPE, "synthetic_data_ack": True},
            headers={"x-csrf-token": self.csrf})
        self.assertEqual(201, response.status_code, response.text)
        return response.json()

    def upload(self, data=None, complete=True, filename="synthetic.xlsx"):
        data = data if data is not None else workbook_bytes()
        intent = self.intent(data, filename)
        response = self.client.put(intent["upload_instruction"]["target"], content=data,
            headers={**intent["upload_instruction"]["headers"], "x-csrf-token": self.csrf})
        self.assertEqual(200, response.status_code, response.text)
        if complete:
            response = self.client.post(f"/api/v1/uploads/{intent['upload_id']}/complete", json={},
                headers={"x-csrf-token": self.csrf})
            self.assertEqual(200, response.status_code, response.text)
        return intent

    def create(self, intent, key=None):
        return self.client.post("/api/v1/reports", json={"upload_id": intent["upload_id"], "summary_mode": "offline"},
            headers={"x-csrf-token": self.csrf, "idempotency-key": key or str(uuid4())})

    def test_valid_workbook_completed_pdf_markdown_excel_match_snapshot(self):
        intent = self.upload()
        response = self.create(intent)
        self.assertEqual(201, response.status_code, response.text)
        report = response.json()
        self.assertEqual("offline", report["summary_source"])
        self.assertEqual("600.0", report["kpis"]["total_budget"]["value"])
        self.assertEqual("0.00%", report["kpis"]["variance_pct"]["display_value"])
        self.assertIs(type(report["kpis"]["transaction_count"]), int)
        self.assertNotIn("object_key", json.dumps(report))
        self.assertNotIn("transactions", report)
        self.assertFalse(self.service.storage.path(self.service.uploads[intent["upload_id"]].object_key).exists())
        identifier = report["report_id"]
        self.assertEqual(report, self.client.get(f"/api/v1/reports/{identifier}").json())
        pdf = self.client.get(f"/api/v1/reports/{identifier}/download/pdf")
        self.assertEqual(200, pdf.status_code)
        self.assertTrue(pdf.content.startswith(b"%PDF-"))
        reader = PdfReader(BytesIO(pdf.content))
        text = " ".join(page.extract_text() for page in reader.pages)
        for expected in ["Laporan Rekap Anggaran", "Rp600", "Rp0", "0.00%", "Total Transaksi", "hypothesis", "Halaman 1", "data sintetis"]:
            self.assertIn(expected, text)
        self.assertNotIn(str(ROOT), text + str(reader.metadata))
        markdown = self.client.get(f"/api/v1/reports/{identifier}/download/markdown").text
        self.assertIn("| Total Budget | Rp600 |", markdown)
        excel = self.client.get(f"/api/v1/reports/{identifier}/download/excel")
        book = load_workbook(BytesIO(excel.content), data_only=True)
        values = dict(book["KPI"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(600, values["Total Budget"])
        self.assertEqual(3, values["Transaction Count"])
        book.close()

    def test_invalid_schema_safe_failed_report_and_replay(self):
        intent = self.upload(workbook_bytes(invalid=True))
        key = str(uuid4())
        failed = self.create(intent, key)
        self.assertEqual(422, failed.status_code)
        error = failed.json()["error"]
        self.assertEqual("INVALID_FILE_SCHEMA", error["code"])
        self.assertIn("report_id", error)
        self.assertNotIn("unknown", failed.text)
        self.assertEqual(error["report_id"], self.create(intent, key).json()["error"]["report_id"])
        self.assertEqual("failed", self.client.get(f"/api/v1/reports/{error['report_id']}").json()["state"])
        self.assertFalse(any(self.service.storage.root.iterdir()))
        self.assertEqual(409, self.create(intent).status_code)

    def test_bad_workbook_security_and_package_limits(self):
        valid = workbook_bytes()
        cases = [b"not-a-zip", workbook_bytes(formula=True),
            modify_archive(valid, extra={"xl/vbaProject.bin": b"macro"}),
            modify_archive(valid, extra={"xl/connections.xml": b"<connections/>"}),
            modify_archive(valid, extra={"xl/_rels/test.rels": b'<Relationships><Relationship TargetMode="External" Target="https://example.invalid"/></Relationships>'}),
            modify_archive(valid, extra={f"padding/{index}.txt": b"x" for index in range(130)}),
            modify_archive(valid, extra={"padding.txt": b"x" * (21 * 1024 * 1024)}),
            modify_archive(valid, replacements={"xl/worksheets/sheet1.xml": lambda value: value.replace(b'A1:E4', b'A1:XFD1048576')}),
            modify_archive(valid, replacements={"xl/worksheets/sheet1.xml": lambda value: value.replace(b'A1:E4', b'A1:E10002')}),
            modify_archive(valid, replacements={"xl/worksheets/sheet1.xml": lambda value: value.replace(b'r="A2"', b'r="A1"')}),
            modify_archive(valid, extra={"xl/worksheets/sheet2.xml": b'<worksheet><sheetData/></worksheet>'}),
            modify_archive(valid, replacements={"xl/workbook.xml": lambda value: value.replace(b'<workbookProtection', b'<workbookProtection lockStructure="1"')}),
            bytes.fromhex("D0CF11E0A1B11AE1") + b"fake-encrypted-workbook",
        ]
        for data in cases:
            with self.subTest(size=len(data)):
                intent = self.upload(data, complete=False)
                response = self.client.post(f"/api/v1/uploads/{intent['upload_id']}/complete", json={}, headers={"x-csrf-token": self.csrf})
                self.assertEqual(422, response.status_code, response.text)
                self.assertNotIn("report_id", response.json()["error"])
                self.assertFalse(self.service.storage.path(self.service.uploads[intent["upload_id"]].object_key).exists())

    def test_oversized_and_wrong_extension_are_rejected(self):
        body = {"filename": "large.xlsx", "size_bytes": MAX_UPLOAD_BYTES + 1, "content_type": XLSX_TYPE, "synthetic_data_ack": True}
        response = self.client.post("/api/v1/uploads/intents", json=body, headers={"x-csrf-token": self.csrf})
        self.assertEqual(413, response.status_code)
        body.update(filename="macro.xlsm", size_bytes=100)
        self.assertEqual(415, self.client.post("/api/v1/uploads/intents", json=body, headers={"x-csrf-token": self.csrf}).status_code)

    def test_actual_upload_stream_is_bounded_and_partial_file_deleted(self):
        body = {"filename": "large.xlsx", "size_bytes": MAX_UPLOAD_BYTES, "content_type": XLSX_TYPE, "synthetic_data_ack": True}
        intent = self.client.post("/api/v1/uploads/intents", json=body, headers={"x-csrf-token": self.csrf}).json()
        response = self.client.put(intent["upload_instruction"]["target"], content=b"x" * (MAX_UPLOAD_BYTES + 1),
            headers={**intent["upload_instruction"]["headers"], "x-csrf-token": self.csrf})
        self.assertEqual(413, response.status_code)
        self.assertFalse(any(self.service.storage.root.iterdir()))

    def test_capability_expiry_reuse_and_size_mismatch(self):
        data = workbook_bytes()
        intent = self.intent(data)
        self.now += timedelta(seconds=121)
        response = self.client.put(intent["upload_instruction"]["target"], content=data,
            headers={**intent["upload_instruction"]["headers"], "x-csrf-token": self.csrf})
        self.assertEqual(410, response.status_code)
        intent = self.upload(complete=False)
        self.assertEqual(403, self.client.put(intent["upload_instruction"]["target"], content=data,
            headers={**intent["upload_instruction"]["headers"], "x-csrf-token": self.csrf}).status_code)
        intent = self.intent(data)
        mismatch = self.client.put(intent["upload_instruction"]["target"], content=data[:-1],
            headers={**intent["upload_instruction"]["headers"], "x-csrf-token": self.csrf})
        self.assertEqual(422, mismatch.status_code)

    def test_cross_owner_upload_report_download_are_not_found(self):
        intent = self.intent(workbook_bytes())
        other = self.new_client()
        other_csrf = self.bootstrap(other)
        self.assertEqual(404, other.put(intent["upload_instruction"]["target"], content=workbook_bytes(),
            headers={**intent["upload_instruction"]["headers"], "x-csrf-token": other_csrf}).status_code)
        self.assertEqual(404, other.post(f"/api/v1/uploads/{intent['upload_id']}/complete", json={}, headers={"x-csrf-token": other_csrf}).status_code)
        self.assertEqual(404, other.post("/api/v1/reports", json={"upload_id": intent["upload_id"]},
            headers={"x-csrf-token": other_csrf, "idempotency-key": str(uuid4())}).status_code)
        report = self.create(self.upload()).json()
        for suffix in ["", "/download/pdf"]:
            self.assertEqual(404, other.get(f"/api/v1/reports/{report['report_id']}{suffix}").status_code)
        other.close()

    def test_csrf_and_synthetic_ack_required(self):
        response = self.client.post("/api/v1/uploads/intents", json={"filename": "x.xlsx", "size_bytes": 12, "content_type": XLSX_TYPE, "synthetic_data_ack": True})
        self.assertEqual(403, response.status_code)
        response = self.client.post("/api/v1/uploads/intents", json={"filename": "x.xlsx", "size_bytes": 12, "content_type": XLSX_TYPE, "synthetic_data_ack": False}, headers={"x-csrf-token": self.csrf})
        self.assertEqual(422, response.status_code)

    def test_idempotency_replay_conflict_and_consumed_upload(self):
        intent, key = self.upload(), str(uuid4())
        first = self.create(intent, key)
        self.assertEqual(201, first.status_code, first.text)
        self.assertEqual(first.json(), self.create(intent, key).json())
        second = self.upload()
        self.assertEqual(409, self.create(second, key).status_code)
        self.assertEqual(409, self.create(intent).status_code)
        self.assertEqual(422, self.create(second, "short").status_code)

    def test_concurrent_duplicate_returns_report_in_progress(self):
        intent = self.upload()
        owner = self.service.uploads[intent["upload_id"]].owner
        entered, release = Event(), Event()
        from mining_api.report_service import analyze_workbook as actual_analyze
        def slow(path):
            entered.set()
            if not release.wait(5):
                raise RuntimeError("test coordination timed out")
            return actual_analyze(path)
        results = []
        key = str(uuid4())
        with patch("mining_api.report_service.analyze_workbook", side_effect=slow):
            thread = Thread(target=lambda: results.append(self.service.create_report(owner, intent["upload_id"], key)))
            thread.start()
            try:
                self.assertTrue(entered.wait(5))
                with self.assertRaises(ApiError) as error:
                    self.service.create_report(owner, intent["upload_id"], key)
                self.assertEqual("REPORT_IN_PROGRESS", error.exception.code)
            finally:
                release.set(); thread.join(5)
        self.assertEqual(1, len(results))

    def test_processing_failure_cleans_partial_artifacts(self):
        intent = self.upload()
        def fail(report, destination):
            destination.write_bytes(b"partial-pdf")
            raise RuntimeError("private-provider-message")
        with patch("mining_api.report_service.render_pdf", side_effect=fail):
            response = self.create(intent)
        self.assertEqual(500, response.status_code)
        self.assertNotIn("private-provider-message", response.text)
        self.assertIn("report_id", response.json()["error"])
        self.assertFalse(any(self.service.storage.root.iterdir()))

    def test_expiry_cleanup_and_filename_traversal_cannot_escape(self):
        intent = self.upload(filename="../../outside.xlsx")
        upload = self.service.uploads[intent["upload_id"]]
        self.assertEqual(self.service.storage.root, self.service.storage.path(upload.object_key).parent)
        report = self.create(intent).json()
        self.now += timedelta(seconds=3601)
        self.service.cleanup()
        self.assertFalse(any(self.service.storage.root.iterdir()))
        self.assertEqual(0, len(self.service.replays))
        self.assertEqual(0, len(self.service.uploads))
        with self.assertRaises(ValueError):
            self.service.storage.path("../../outside")

    def test_pdf_long_tables_escape_markup_and_null_percentage(self):
        report = self.create(self.upload()).json()
        report["kpis"]["variance_pct"] = {"value": None, "display_value": "N/A (budget nol)"}
        row = report["category_summary"][0]
        report["category_summary"] = [dict(deepcopy(row), category=f"Category {i} <img src='https://example.invalid'>") for i in range(100)]
        report["summary"]["executive_summary"] = "Teks panjang & aman <b>literal</b>. " * 100
        destination = Path(self.temp.name) / "long.pdf"
        render_pdf(report, destination)
        reader = PdfReader(destination)
        self.assertGreater(len(reader.pages), 3)
        text = " ".join(page.extract_text() for page in reader.pages)
        self.assertIn("N/A (budget nol)", text)
        self.assertIn("<img", text)
        self.assertIn("<b>literal</b>", text)
        self.assertGreater(sum("Nama" in page.extract_text() for page in reader.pages), 2)

    def test_baseline_pdf_all_kpi_and_group_display_values_match(self):
        report = self.create(self.upload((ROOT / "data/raw/accounting_sample.xlsx").read_bytes())).json()
        pdf = self.client.get(f"/api/v1/reports/{report['report_id']}/download/pdf")
        text = " ".join(page.extract_text() for page in PdfReader(BytesIO(pdf.content)).pages)
        self.assertEqual(120, report["kpis"]["transaction_count"])
        self.assertEqual(60, report["kpis"]["over_budget_count"])
        self.assertEqual(60, report["kpis"]["under_budget_count"])
        for section in (report["department_summary"], report["category_summary"]):
            for row in section:
                for key in ("budget", "actual_cost", "variance", "variance_pct"):
                    self.assertIsInstance(row[key]["value"], (str, type(None)))
                    self.assertIn(row[key]["display_value"], text)
        self.assertIn("Rp902.850.000", text)
        self.assertIn("2.50%", text)
        self.assertIn("Rp589.000.000", text)
        self.assertIn("7.67%", text)

    def test_pdf_embeds_fonts_and_uses_indonesian_headings(self):
        report = self.create(self.upload()).json()
        original = deepcopy(report)
        destination = Path(self.temp.name) / "embedded.pdf"
        render_pdf(report, destination)
        self.assertEqual(original, report)
        reader = PdfReader(destination)
        text = " ".join(page.extract_text() for page in reader.pages)
        for heading in ["Ringkasan Eksekutif", "Temuan Utama", "Potensi Risiko", "Rekomendasi Tindakan"]:
            self.assertIn(heading, text)
        embedded = set()
        for page in reader.pages:
            for reference in page["/Resources"]["/Font"].values():
                font = reference.get_object()
                descriptor = font.get("/FontDescriptor")
                if descriptor is not None and "/FontFile2" in descriptor.get_object():
                    embedded.add(str(font["/BaseFont"]))
        self.assertTrue(any("Vera" in name for name in embedded), embedded)
        self.assertTrue(any("Bold" in name for name in embedded), embedded)

    def test_zero_budget_pipeline_serializes_null_and_pdf_na(self):
        book = load_workbook(BytesIO(workbook_bytes()))
        for row in range(2, 5):
            book.active.cell(row, 4, 0)
        output = BytesIO(); book.save(output); book.close()
        report = self.create(self.upload(output.getvalue())).json()
        self.assertIsNone(report["kpis"]["variance_pct"]["value"])
        pdf = self.client.get(f"/api/v1/reports/{report['report_id']}/download/pdf")
        self.assertIn("N/A (budget nol)", " ".join(page.extract_text() for page in PdfReader(BytesIO(pdf.content)).pages))

    def test_cleanup_failure_is_recorded_and_retried(self):
        intent = self.upload(workbook_bytes(invalid=True))
        with patch.object(self.service.storage, "delete", side_effect=OSError("fake cleanup failure")):
            response = self.create(intent)
        self.assertEqual(422, response.status_code)
        self.assertEqual(1, len(self.service.pending_cleanup))
        self.service.cleanup()
        self.assertFalse(self.service.pending_cleanup)
        self.assertFalse(any(self.service.storage.root.iterdir()))

    def test_session_reset_never_adopts_old_owner(self):
        report = self.create(self.upload()).json()
        response = self.client.post("/api/v1/session/reset", json={"reset_ack": True})
        self.assertEqual(201, response.status_code)
        self.assertEqual(404, self.client.get(f"/api/v1/reports/{report['report_id']}").status_code)


if __name__ == "__main__":
    unittest.main()
