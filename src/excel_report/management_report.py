"""Markdown management report rendering from deterministic KPI output."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook

from .models import ExecutiveSummaryResult, ReportExportError


def _format_idr(value: float | int) -> str:
    return "Rp" + f"{value:,.0f}".replace(",", ".")


def _format_percentage(value: float | None) -> str:
    return "N/A" if value is None else f"{value:.2f}%"


def _summary_table(rows: list[dict[str, Any]], label_key: str) -> list[str]:
    lines = [
        "| Nama | Budget | Actual Cost | Variance | Variance % | Transaksi |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for row in rows:
        lines.append(
            f"| {row[label_key]} | {_format_idr(row['budget'])} | "
            f"{_format_idr(row['actual_cost'])} | {_format_idr(row['variance'])} | "
            f"{_format_percentage(row['variance_pct'])} | {row['transaction_count']} |"
        )
    return lines


def render_management_report(
    payload: dict[str, Any],
    summary: ExecutiveSummaryResult,
) -> str:
    """Render a management report while sourcing every KPI from Python output."""
    metadata = payload["metadata"]
    kpis = payload["overall_kpis"]
    source_labels = {
        "offline": "Offline deterministic template",
        "gemini": "Gemini structured interpretation",
        "offline_fallback": "Offline deterministic fallback",
    }
    lines = [
        "# Mining Operational Cost Management Report",
        "",
        f"> **Disclaimer:** {summary.synthetic_data_disclaimer}",
        "",
        "## Metadata",
        "",
        f"- Periode: {metadata['period_start']} sampai {metadata['period_end']}",
        f"- Currency: {metadata['currency']}",
        f"- Synthetic data: {str(metadata['synthetic_data']).lower()}",
        f"- Summary source: {source_labels[summary.source]}",
    ]
    if summary.fallback_reason:
        lines.append(f"- Fallback status: {summary.fallback_reason}")
    lines.extend(
        [
            "",
            "## KPI Utama",
            "",
            "| KPI | Nilai |",
            "|---|---:|",
            f"| Total Budget | {_format_idr(kpis['total_budget'])} |",
            f"| Total Actual Cost | {_format_idr(kpis['total_actual_cost'])} |",
            f"| Variance | {_format_idr(kpis['variance'])} |",
            f"| Variance % | {_format_percentage(kpis['variance_pct'])} |",
            f"| Total Transaksi | {kpis['transaction_count']} |",
            f"| Over-budget | {kpis['over_budget_count']} |",
            f"| Under-budget | {kpis['under_budget_count']} |",
            f"| On-budget | {kpis['on_budget_count']} |",
            "",
            "Konvensi: `variance = actual_cost - budget`; variance positif berarti "
            "over-budget.",
            "",
            "## Ringkasan Departemen",
            "",
            *_summary_table(payload["department_summary"], "department"),
            "",
            "## Ringkasan Kategori",
            "",
            *_summary_table(payload["category_summary"], "category"),
            "",
            "## Executive Summary",
            "",
            summary.executive_summary,
            "",
            "## Key Findings",
            "",
            *[f"- {finding}" for finding in summary.key_findings],
            "",
            "## Potential Risks",
            "",
            *[
                f"- **{risk['evidence_level'].capitalize()}:** {risk['statement']}"
                for risk in summary.potential_risks
            ],
            "",
            "## Actionable Recommendations",
            "",
            *[
                f"{index}. {recommendation}"
                for index, recommendation in enumerate(
                    summary.actionable_recommendations, start=1
                )
            ],
            "",
            "---",
            "Angka pada tabel berasal langsung dari KPI engine Python. Interpretasi tidak "
            "mengubah nilai tersebut.",
            "",
        ]
    )
    return "\n".join(lines)


def export_management_report(
    payload: dict[str, Any],
    summary: ExecutiveSummaryResult,
    output_path: Path,
) -> Path:
    """Write the final Markdown management report."""
    destination = Path(output_path)
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text(
            render_management_report(payload, summary),
            encoding="utf-8",
        )
    except OSError as error:
        raise ReportExportError(
            f"Unable to write management report '{destination}': {error}"
        ) from error
    return destination


def export_management_workbook(
    payload: dict[str, Any],
    summary: ExecutiveSummaryResult,
    output_path: Path,
) -> Path:
    """Write numeric KPIs, aggregates, and the validated summary to Excel."""
    destination = Path(output_path)
    workbook = Workbook()
    kpi_sheet = workbook.active
    kpi_sheet.title = "KPI"
    kpis = payload["overall_kpis"]
    kpi_sheet.append(["Metric", "Value"])
    for label, key in (
        ("Total Budget", "total_budget"),
        ("Total Actual Cost", "total_actual_cost"),
        ("Variance", "variance"),
        ("Variance %", "variance_pct"),
        ("Transaction Count", "transaction_count"),
        ("Over-budget Count", "over_budget_count"),
        ("Under-budget Count", "under_budget_count"),
        ("On-budget Count", "on_budget_count"),
    ):
        kpi_sheet.append([label, kpis[key]])

    for sheet_name, rows, label_key in (
        ("Department Summary", payload["department_summary"], "department"),
        ("Category Summary", payload["category_summary"], "category"),
    ):
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(
            [
                label_key,
                "budget",
                "actual_cost",
                "variance",
                "variance_pct",
                "transaction_count",
            ]
        )
        for row in rows:
            worksheet.append(
                [
                    row[label_key],
                    row["budget"],
                    row["actual_cost"],
                    row["variance"],
                    row["variance_pct"],
                    row["transaction_count"],
                ]
            )

    summary_sheet = workbook.create_sheet("Executive Summary")
    summary_sheet.append(["Summary Source", summary.source])
    summary_sheet.append(["Synthetic Data Disclaimer", summary.synthetic_data_disclaimer])
    summary_sheet.append(["Executive Summary", summary.executive_summary])
    if summary.fallback_reason:
        summary_sheet.append(["Fallback Reason", summary.fallback_reason])
    for finding in summary.key_findings:
        summary_sheet.append(["Key Finding", finding])
    for risk in summary.potential_risks:
        summary_sheet.append(
            [f"Potential Risk ({risk['evidence_level']})", risk["statement"]]
        )
    for recommendation in summary.actionable_recommendations:
        summary_sheet.append(["Actionable Recommendation", recommendation])

    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
    except (OSError, ValueError) as error:
        raise ReportExportError(
            f"Unable to write management workbook '{destination}': {error}"
        ) from error
    finally:
        workbook.close()
    return destination
