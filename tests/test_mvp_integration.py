"""End-to-end MVP test using only temporary synthetic files."""

from __future__ import annotations

from contextlib import redirect_stdout
from io import StringIO
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
import pandas as pd


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
SRC_PATH = REPOSITORY_ROOT / "src"
if str(SRC_PATH) not in sys.path:
    sys.path.insert(0, str(SRC_PATH))

from analyze_costs import main  # noqa: E402


class MvpIntegrationTests(unittest.TestCase):
    def test_offline_cli_keeps_kpis_consistent_across_outputs(self) -> None:
        dataframe = pd.DataFrame(
            {
                "date": pd.to_datetime(["2025-01-01", "2025-01-02", "2025-01-03"]),
                "department": ["Mining Operations", "HSE", "Logistics"],
                "category": ["Diesel Fuel", "Safety Equipment", "Hauling Transportation"],
                "budget": [100.0, 200.0, 300.0],
                "actual_cost": [120.0, 180.0, 300.0],
            }
        )

        with tempfile.TemporaryDirectory() as temporary_directory:
            temporary_path = Path(temporary_directory)
            input_path = temporary_path / "input.xlsx"
            output_dir = temporary_path / "processed"
            markdown_path = temporary_path / "management_report.md"
            workbook_path = temporary_path / "management_report.xlsx"
            dataframe.to_excel(input_path, index=False, engine="openpyxl")
            arguments = [
                "analyze_costs.py",
                "--input",
                str(input_path),
                "--output-dir",
                str(output_dir),
                "--report-path",
                str(markdown_path),
                "--excel-report-path",
                str(workbook_path),
                "--mode",
                "offline",
            ]

            with patch.object(sys, "argv", arguments), redirect_stdout(StringIO()):
                exit_code = main()

            self.assertEqual(0, exit_code)
            with (output_dir / "analysis_summary.json").open(encoding="utf-8") as file:
                payload = json.load(file)
            markdown = markdown_path.read_text(encoding="utf-8")
            workbook = load_workbook(workbook_path, data_only=True)
            try:
                kpi_values = {
                    row[0]: row[1]
                    for row in workbook["KPI"].iter_rows(
                        min_row=2, values_only=True
                    )
                }
                department_rows = list(
                    workbook["Department Summary"].iter_rows(values_only=True)
                )
                category_rows = list(
                    workbook["Category Summary"].iter_rows(values_only=True)
                )
            finally:
                workbook.close()

            expected_kpis = {
                "total_budget": 600.0,
                "total_actual_cost": 600.0,
                "variance": 0.0,
                "variance_pct": 0.0,
                "transaction_count": 3,
                "over_budget_count": 1,
                "under_budget_count": 1,
                "on_budget_count": 1,
            }
            self.assertEqual(expected_kpis, payload["overall_kpis"])
            self.assertIn("| Total Budget | Rp600 |", markdown)
            self.assertIn("| Total Actual Cost | Rp600 |", markdown)
            self.assertIn("| Variance | Rp0 |", markdown)
            self.assertIn("| Variance % | 0.00% |", markdown)
            self.assertIn("| Total Transaksi | 3 |", markdown)
            self.assertIn("| Over-budget | 1 |", markdown)
            self.assertIn("| Under-budget | 1 |", markdown)
            self.assertIn("| On-budget | 1 |", markdown)
            self.assertEqual(600, kpi_values["Total Budget"])
            self.assertEqual(600, kpi_values["Total Actual Cost"])
            self.assertEqual(0, kpi_values["Variance"])
            self.assertEqual(0, kpi_values["Variance %"])
            self.assertEqual(3, kpi_values["Transaction Count"])
            self.assertEqual(1, kpi_values["Over-budget Count"])
            self.assertEqual(1, kpi_values["Under-budget Count"])
            self.assertEqual(1, kpi_values["On-budget Count"])

            department_header = department_rows[0]
            department_values = dict(zip(department_header, department_rows[1]))
            self.assertEqual("Mining Operations", department_values["department"])
            self.assertEqual(100, department_values["budget"])
            self.assertEqual(120, department_values["actual_cost"])
            self.assertEqual(20, department_values["variance"])
            self.assertEqual(20, department_values["variance_pct"])
            self.assertEqual(1, department_values["transaction_count"])
            self.assertIn(
                "| Mining Operations | Rp100 | Rp120 | Rp20 | 20.00% | 1 |",
                markdown,
            )

            category_header = category_rows[0]
            category_values = dict(zip(category_header, category_rows[1]))
            self.assertEqual("Diesel Fuel", category_values["category"])
            self.assertEqual(100, category_values["budget"])
            self.assertEqual(120, category_values["actual_cost"])
            self.assertEqual(20, category_values["variance"])
            self.assertEqual(20, category_values["variance_pct"])
            self.assertEqual(1, category_values["transaction_count"])
            self.assertIn(
                "| Diesel Fuel | Rp100 | Rp120 | Rp20 | 20.00% | 1 |",
                markdown,
            )
            self.assertTrue(markdown_path.is_file())
            self.assertTrue(workbook_path.is_file())


if __name__ == "__main__":
    unittest.main()
